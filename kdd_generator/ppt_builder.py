# kdd_generator/ppt_builder.py — PowerPoint export using python-pptx
# Called from server.py /api/kdd/ppt endpoint
import json
from datetime import date
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

from .helpers import dedup, get_rag_status

ROOT = Path(__file__).resolve().parent.parent

_ACCENTURE_PURPLE = RGBColor(0x82, 0x00, 0xED)
_DARK_BLUE        = RGBColor(0x1F, 0x4E, 0x79)
_WHITE            = RGBColor(0xFF, 0xFF, 0xFF)
_RED_RGB          = RGBColor(0xFF, 0x00, 0x00)
_AMBER_RGB        = RGBColor(0xFF, 0xBF, 0x00)
_GREEN_RGB        = RGBColor(0x00, 0x80, 0x00)

_RAG_COLOR = {"Red": _RED_RGB, "Amber": _AMBER_RGB, "Green": _GREEN_RGB}


def _add_text(tf, text, size=12, bold=False, color=None):
    tf.text = ""
    p  = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    run = p.add_run()
    run.text = str(text or "")
    run.font.size  = Pt(size)
    run.font.bold  = bold
    if color:
        run.font.color.rgb = color


def _cover_slide(prs: Presentation, meta: dict):
    blank = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank)

    # Purple background
    bg   = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = _ACCENTURE_PURPLE

    def _txbox(left, top, width, height, text, size=24, bold=False, color=_WHITE):
        tb = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
        tf = tb.text_frame
        tf.word_wrap = True
        _add_text(tf, text, size=size, bold=bold, color=color)
        return tb

    today = date.today().strftime("%d %B %Y")
    _txbox(0.8, 0.8,  8.4, 0.8, "ACCENTURE", size=32, bold=True)
    _txbox(0.8, 1.8,  8.4, 1.0, "Key Design Decisions (KDD) Log", size=28, bold=True)
    _txbox(0.8, 2.9,  8.4, 0.6, "SAP S/4HANA Cloud Public Edition", size=16)
    _txbox(0.8, 3.7,  8.4, 0.5, f"Client:  {meta.get('client', '[Client Name]')}", size=14)
    _txbox(0.8, 4.3,  8.4, 0.5, f"Project: {meta.get('project', '[Project Name]')}", size=14)
    _txbox(0.8, 4.9,  8.4, 0.5, f"Generated: {today}", size=12)
    _txbox(0.8, 5.5,  8.4, 0.5, f"Total KDD Items: {meta.get('totalKDDs', '')}", size=12)


def _summary_slide(prs: Presentation, rows: list):
    blank = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank)
    bg = slide.background.fill
    bg.solid(); bg.fore_color.rgb = _DARK_BLUE

    def _lbl(left, top, w, h, text, size=11, bold=False, color=_WHITE):
        tb = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(w), Inches(h))
        tf = tb.text_frame
        tf.word_wrap = True
        _add_text(tf, text, size=size, bold=bold, color=color)

    _lbl(0.5, 0.2, 9, 0.6, "KDD Summary", size=20, bold=True)
    fit   = sum(1 for r in rows if r.get("fitGap") == "Fit")
    part  = sum(1 for r in rows if r.get("fitGap") == "Partial Fit")
    gap   = sum(1 for r in rows if r.get("fitGap") == "Gap")
    red   = sum(1 for r in rows if get_rag_status(r) == "Red")
    amber = sum(1 for r in rows if get_rag_status(r) == "Amber")
    green = sum(1 for r in rows if get_rag_status(r) == "Green")
    ricef = sum(1 for r in rows if r.get("ricefType") and r["ricefType"] != "None")

    _lbl(0.5, 1.0, 4, 0.4, f"Total KDD Items: {len(rows)}", size=14, bold=True)
    _lbl(0.5, 1.5, 4, 0.4, f"Fit: {fit}  |  Partial Fit: {part}  |  Gap: {gap}")
    _lbl(0.5, 2.0, 4, 0.4, f"Red: {red}  |  Amber: {amber}  |  Green: {green}")
    _lbl(0.5, 2.5, 4, 0.4, f"RICEF Items: {ricef}")


def _kdd_slide(prs: Presentation, row: dict, index: int, total: int):
    blank = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank)

    rag   = get_rag_status(row)
    color = _RAG_COLOR.get(rag, _GREEN_RGB)

    def _box(left, top, w, h, text, size=10, bold=False, color=None):
        tb = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(w), Inches(h))
        tf = tb.text_frame
        tf.word_wrap = True
        _add_text(tf, text, size=size, bold=bold, color=color)

    # Header bar
    rect = slide.shapes.add_shape(1, Inches(0), Inches(0), Inches(10), Inches(0.55))
    rect.fill.solid(); rect.fill.fore_color.rgb = _DARK_BLUE
    rect.line.fill.background()

    _box(0.15, 0.08, 6, 0.4, f"[{index}/{total}] {row.get('kddId', '')} — {row.get('scopeItemName', '')}", size=11, bold=True, color=_WHITE)
    _box(7.5,  0.08, 2, 0.4, f"RAG: {rag}", size=11, bold=True, color=color)

    _box(0.15, 0.7,  9.7, 0.5, row.get("designQuestion", ""), size=12, bold=True)
    _box(0.15, 1.3,  2,   0.35, f"Fit/Gap: {row.get('fitGap', '')}", size=10)
    _box(2.5,  1.3,  2.5, 0.35, f"Complexity: {row.get('complexity', '')}", size=10)
    _box(5.5,  1.3,  4,   0.35, f"Phase: {row.get('sapActivatePhase', '')}", size=10)

    _box(0.15, 1.75, 9.7, 0.35, "Rationale:", size=9, bold=True)
    _box(0.15, 2.1,  9.7, 0.7,  row.get("rationale", ""), size=9)
    _box(0.15, 2.85, 9.7, 0.35, "Impact / Actions:", size=9, bold=True)
    _box(0.15, 3.2,  9.7, 0.7,  row.get("impactActions", ""), size=9)
    _box(0.15, 3.95, 9.7, 0.35, f"Decision Owner: {row.get('decisionOwner', '')}  |  Fiori / BTP: {row.get('notes', '')}", size=9)

    if row.get("ricefType") and row["ricefType"] != "None":
        _box(0.15, 4.4, 9.7, 0.35, f"RICEF: {row.get('ricefTitle', '')}  [{row.get('ricefType', '')}]  |  Extensibility: {row.get('extensibilityType', '')}  |  BTP: {row.get('btpRequired', '')}", size=9)


def build(json_path: Path) -> Path:
    data     = json.loads(json_path.read_text("utf-8"))
    rows     = dedup(data.get("rows", []))
    meta     = data.get("meta", {})
    meta["totalKDDs"] = len(rows)

    prs = Presentation()
    prs.slide_width  = Inches(10)
    prs.slide_height = Inches(5.625)

    _cover_slide(prs, meta)
    _summary_slide(prs, rows)
    for i, row in enumerate(rows, start=1):
        _kdd_slide(prs, row, i, len(rows))

    out_dir = ROOT / "output"
    out_dir.mkdir(exist_ok=True)

    client  = "".join(c if c.isalnum() else "_" for c in (meta.get("client") or ""))[:20]
    project = "".join(c if c.isalnum() else "_" for c in (meta.get("project") or ""))[:20]
    today   = date.today().strftime("%Y%m%d")
    parts   = ["Accenture_KDD"]
    if client:  parts.append(client)
    if project: parts.append(project)
    parts.append(today)
    filename = "_".join(parts) + ".pptx"
    out_path = out_dir / filename
    prs.save(str(out_path))
    return out_path


if __name__ == "__main__":
    import sys
    if len(sys.argv) < 2:
        print("Usage: python -m kdd_generator.ppt_builder output/kdd-data-{session}.json", file=sys.stderr)
        sys.exit(1)
    src = Path(sys.argv[1])
    out = build(src)
    print(f"PowerPoint written: {out}")
