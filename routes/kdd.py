# routes/kdd.py — all /api/kdd/* endpoints
import asyncio
import io
import json
import time
from datetime import date, datetime
from typing import AsyncGenerator

from fastapi import APIRouter, HTTPException, Request
from fastapi.responses import StreamingResponse

from .deps import (
    CACHE_PATH, append_run_log, get_past_decisions, kdd_data_path, load_cache, load_catalog,
    slugify, spawn_claude, sse,
)

router = APIRouter(prefix="/api/kdd")


@router.post("/generate")
async def kdd_generate(request: Request):
    body            = await request.json()
    scope_ids       = [s.strip().upper() for s in (body.get("scopeIds") or []) if s.strip()]
    client          = (body.get("client") or "").strip()
    project         = (body.get("project") or "").strip()
    mode            = body.get("mode", "ai")
    client_country  = (body.get("clientCountry") or "").strip()
    client_region   = (body.get("clientRegion") or "").strip()
    start_time = time.monotonic()

    if not scope_ids:
        raise HTTPException(400, "scopeIds required")

    cat     = load_catalog()
    cat_map = {(i.get("id") or "").upper(): i for i in cat.get("items", [])}
    items     = [cat_map[sid] for sid in scope_ids if sid in cat_map]
    not_found = [sid for sid in scope_ids if sid not in cat_map]

    async def generate_stream() -> AsyncGenerator[str, None]:
        if not items:
            yield sse({"type": "fatal", "message": f"Scope IDs not found in catalog: {', '.join(not_found)}"})
            return
        if not_found:
            yield sse({"type": "warning", "message": f"Not found (skipped): {', '.join(not_found)}"})

        cache = load_cache()
        yield sse({"type": "progress", "message": f"Starting {'cache' if mode == 'fast' else 'AI'} pipeline for {len(items)} scope item(s)…"})

        from agents_py import orchestrator
        from kdd_generator.excel_builder import build_excel

        events_q: asyncio.Queue = asyncio.Queue()

        def on_event(ev: dict):
            events_q.put_nowait(ev)

        async def _run():
            try:
                result = await orchestrator.run(items=items, cache_data=cache, on_event=on_event, concurrency=3, mode=mode)
            except Exception as exc:
                events_q.put_nowait({"type": "error", "message": f"Pipeline failed: {exc}"})
                result = {"rows": [], "overviews": [], "src_counts": {"claude": 0, "cache": 0, "error": len(items)}}
            events_q.put_nowait({"type": "_pipeline_done_", "result": result})

        _task = asyncio.create_task(_run())

        all_rows = []
        overviews = []
        src_counts = {}

        try:
            while True:
                try:
                    ev = await asyncio.wait_for(events_q.get(), timeout=0.5)
                except asyncio.TimeoutError:
                    yield "data: \n\n"
                    continue
                if ev.get("type") == "_pipeline_done_":
                    result     = ev["result"]
                    all_rows   = result["rows"]
                    overviews  = result["overviews"]
                    src_counts = result["src_counts"]
                    break
                yield sse(ev)

            if not all_rows:
                yield sse({"type": "fatal", "message": "No KDDs generated — check scope IDs exist in the catalog."})
                return

            yield sse({"type": "progress", "message": "Building Excel workbook…"})
        finally:
            if not _task.done():
                _task.cancel()
                try:
                    await _task
                except (asyncio.CancelledError, Exception):
                    pass
        meta = {
            "client": client, "project": project,
            "catalogVersion": str(cat.get("version", "")),
            "catalogCountry": cat.get("country", "Global") or "Global",
            "clientCountry":  client_country or "Global",
            "clientRegion":   client_region or "",
            "generatedAt":    datetime.utcnow().isoformat() + "Z",
            "scopeCount":     len(overviews),
            "overviews":      overviews,
            "agentPipeline":  True,
        }
        session_id = f"{slugify(client)}-{slugify(project)}" if (client or project) else "default"
        json_path  = kdd_data_path(session_id)
        json_path.write_text(json.dumps({"rows": all_rows, "meta": meta, "sessionId": session_id}, indent=2, ensure_ascii=False), "utf-8")

        # Update kdd-cache.json so cacheReady=true on the status endpoint
        try:
            from collections import defaultdict
            existing_cache = json.loads(CACHE_PATH.read_text("utf-8")) if CACHE_PATH.exists() else {"items": {}}
            cache_items    = existing_cache.get("items", {})
            rows_by_scope: dict = defaultdict(list)
            ov_by_scope    = {ov["id"]: ov.get("overview", "") for ov in overviews}
            for row in all_rows:
                sid = row.get("scopeItemId", "")
                if sid:
                    rows_by_scope[sid].append(row)
            for sid, scope_rows in rows_by_scope.items():
                cache_items[sid] = {
                    "rows":     scope_rows,
                    "overview": ov_by_scope.get(sid, ""),
                    "cachedAt": datetime.utcnow().isoformat() + "Z",
                }
            CACHE_PATH.write_text(
                json.dumps({"items": cache_items, "version": str(cat.get("version", ""))}, indent=2, ensure_ascii=False),
                "utf-8",
            )
        except Exception:
            pass  # cache write is non-fatal

        try:
            out_path    = build_excel(all_rows, meta)
            filename    = out_path.name
            duration_ms = int((time.monotonic() - start_time) * 1000)
            append_run_log({
                "ts": datetime.utcnow().isoformat() + "Z",
                "client": client, "project": project,
                "scopeIds": scope_ids, "status": "ok",
                "filename": filename, "totalKDDs": len(all_rows),
                "durationMs": duration_ms, "sources": src_counts, "agentPipeline": True,
            })
            yield sse({"type": "complete", "filename": filename, "totalKDDs": len(all_rows), "sources": src_counts, "sessionId": session_id})
        except Exception as exc:
            append_run_log({"ts": datetime.utcnow().isoformat() + "Z", "client": client, "project": project,
                            "scopeIds": scope_ids, "status": "excel_error", "totalKDDs": len(all_rows),
                            "durationMs": int((time.monotonic() - start_time) * 1000), "sources": src_counts})
            yield sse({"type": "fatal", "message": f"Excel build failed: {exc}"})

    return StreamingResponse(generate_stream(), media_type="text/event-stream",
                             headers={"Cache-Control": "no-cache", "X-Accel-Buffering": "no"})


@router.post("/suggest")
async def kdd_suggest(request: Request):
    body            = await request.json()
    design_question = body.get("designQuestion", "")
    scope_item_id   = body.get("scopeItemId", "")
    scope_item_name = body.get("scopeItemName", "")
    fit_gap         = body.get("fitGap", "?")
    notes           = body.get("notes", "")

    if not design_question:
        raise HTTPException(400, "designQuestion required")

    past = get_past_decisions(scope_item_id)
    if len(past) >= 3:
        return {"ok": True, "source": "library", "libraryCount": len(past), "suggestions": [{
            "text":      d.get("decision", ""),
            "rationale": d.get("rationale") or f"Based on decision made for {d.get('client', 'a previous project')}",
            "tag":       "Custom Path" if d.get("fitGap") == "Gap" else "Common Variation" if d.get("fitGap") == "Partial Fit" else "SAP Standard",
            "source":    "library",
        } for d in past[:3]]}

    past_block = ""
    if past:
        past_block = f"\nPast decisions for {scope_item_id}:\n" + "\n".join(
            f"• {d.get('decision', '')} [{d.get('fitGap', '?')}]" for d in past
        )

    prompt = f"""SAP S/4HANA Cloud Public Edition consultant. Give exactly 3 decision options.

Question: {design_question}
Scope: {scope_item_id}{' — ' + scope_item_name if scope_item_name else ''}
Fit/Gap: {fit_gap}{chr(10) + 'Fiori app: ' + notes if notes else ''}{past_block}

Rules: Cloud PE only. Option 1=SAP standard, Option 2=client variation, Option 3=extensibility (Key User/BTP). Name specific Fiori app or BTP service. No ABAP/ECC/on-premise/SAP GUI.
IMPORTANT: If you cannot suggest a specific real Fiori app or BTP service, write "To be confirmed in workshop" — never fabricate.

JSON only — no markdown:
[{{"text":"decision (2 sentences max)","rationale":"why (1 sentence)","tag":"SAP Standard|Common Variation|Custom Path"}}]"""

    try:
        out = await spawn_claude(prompt, timeout=90)
        s, e = out.find("["), out.rfind("]")
        if s < 0 or e < 0:
            raise ValueError("No JSON array in response")
        return {"ok": True, "suggestions": json.loads(out[s:e+1]), "source": "claude"}
    except TimeoutError:
        raise HTTPException(504, "Claude timed out — try again")
    except Exception as exc:
        raise HTTPException(500, f"Suggest failed: {exc}")


@router.get("/sessions")
async def kdd_sessions():
    """List all saved KDD sessions sorted by most recent, so the editor can auto-pick the latest."""
    from .deps import OUTPUT_DIR
    files = sorted(
        [f for f in OUTPUT_DIR.glob("kdd-data-*.json")],
        key=lambda f: f.stat().st_mtime, reverse=True,
    )
    sessions = []
    for f in files:
        sid = f.stem.replace("kdd-data-", "", 1)
        try:
            data       = json.loads(f.read_text("utf-8"))
            meta       = data.get("meta", {})
            total_kdds = len(data.get("rows", []))
        except Exception:
            meta       = {}
            total_kdds = 0
        sessions.append({
            "sessionId":   sid,
            "client":      meta.get("client", ""),
            "project":     meta.get("project", ""),
            "generatedAt": meta.get("generatedAt", ""),
            "totalKDDs":   total_kdds,
        })
    return {"sessions": sessions}


@router.get("/data")
async def kdd_data(session_id: str = "default"):
    jp = kdd_data_path(session_id)
    if not jp.exists():
        raise HTTPException(404, "No KDD data yet. Generate KDDs first.")
    return json.loads(jp.read_text("utf-8"))


@router.get("/editor")
async def kdd_editor_get(session_id: str = "default"):
    jp = kdd_data_path(session_id)
    if not jp.exists():
        raise HTTPException(404, "No KDD data yet. Generate KDDs first.")
    data = json.loads(jp.read_text("utf-8"))
    return {"ok": True, "rows": data.get("rows", []), "meta": data.get("meta", {})}


@router.put("/editor")
async def kdd_editor_put(request: Request, session_id: str = "default"):
    from datetime import datetime
    body     = await request.json()
    rows     = body.get("rows", [])
    meta     = body.get("meta", {})
    jp       = kdd_data_path(session_id)
    existing = {}
    if jp.exists():
        try:
            existing = json.loads(jp.read_text("utf-8"))
        except Exception:
            pass
    saved_at = datetime.utcnow().isoformat() + "Z"
    merged_meta = {**existing.get("meta", {}), **meta, "lastEditedAt": saved_at}
    payload = {**existing, "rows": rows, "meta": merged_meta}
    jp.write_text(json.dumps(payload, indent=2, ensure_ascii=False), "utf-8")
    return {"ok": True, "count": len(rows), "savedAt": saved_at}


@router.post("/retry/{kdd_id}")
async def kdd_retry_row(kdd_id: str, request: Request):
    body = await request.json()
    row  = body.get("row", {})
    if not row:
        raise HTTPException(400, "row required")
    scope_id = row.get("scopeItemId", "")
    cat      = load_catalog()
    item     = next((i for i in cat.get("items", []) if (i.get("id") or "").upper() == scope_id.upper()), None)
    if not item:
        raise HTTPException(404, f"Scope item {scope_id} not found in catalog")
    from agents_py import context_agent, generation_agent, validation_agent
    state = {"item": item, "context": None, "past_decisions": get_past_decisions(scope_id), "on_event": None}
    try:
        state = await context_agent.run(state)
        state = await generation_agent.run(state)
        state = await validation_agent.run(state)
    except (TimeoutError, RuntimeError) as exc:
        raise HTTPException(500, f"Retry generation failed: {exc}")
    new_rows = state.get("rows", [])
    matching = next((r for r in new_rows if r.get("kddId") == kdd_id), new_rows[0] if new_rows else None)
    if not matching:
        raise HTTPException(500, "Retry produced no rows")
    return {"ok": True, "row": matching}


@router.post("/export")
async def kdd_export(request: Request, session_id: str = "default"):
    try:
        body = await request.json()
    except Exception:
        body = {}
    jp = kdd_data_path(session_id)
    if not jp.exists():
        raise HTTPException(404, "No KDD data. Generate KDDs first.")
    data = json.loads(jp.read_text("utf-8"))
    rows = data.get("rows", [])
    if not rows:
        raise HTTPException(400, "No KDD rows to export.")
    meta = {**data.get("meta", {}), **{k: v for k, v in body.items() if v}}
    try:
        from kdd_generator.excel_builder import build_excel
        out_path = build_excel(rows, meta)
        return {"ok": True, "filename": out_path.name}
    except Exception as exc:
        raise HTTPException(500, f"Excel build failed: {exc}")


@router.post("/ppt")
async def kdd_ppt(session_id: str = "default"):
    jp = kdd_data_path(session_id)
    if not jp.exists():
        raise HTTPException(404, "No KDD data yet. Generate KDDs first.")
    from kdd_generator import ppt_builder
    out_path = ppt_builder.build(jp)
    return {"ok": True, "filename": out_path.name}


@router.post("/wricef")
async def kdd_wricef(request: Request, session_id: str = "default"):
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    jp = kdd_data_path(session_id)
    if not jp.exists():
        raise HTTPException(404, "No KDD data yet.")
    data  = json.loads(jp.read_text("utf-8"))
    rows  = [r for r in data.get("rows", []) if r.get("ricefType") and r["ricefType"] != "None"]

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "WRICEF Log"
    headers = ["KDD ID","Scope Item ID","Scope Item Name","Design Question","Fit/Gap","RICEF Title",
               "RICEF Type","Extensibility Path","BTP Required","Effort Estimate","Complexity","Decision Owner","Notes"]
    ws.append(headers)
    hfill = PatternFill("solid", fgColor="1F4E79")
    hfont = Font(color="FFFFFF", bold=True)
    for cell in ws[1]:
        cell.font = hfont
        cell.fill = hfill
    for r in rows:
        ws.append([
            r.get("kddId",""), r.get("scopeItemId",""), r.get("scopeItemName",""),
            r.get("designQuestion",""), r.get("fitGap",""), r.get("ricefTitle",""),
            r.get("ricefType",""), r.get("extensibilityType",""), r.get("btpRequired",""),
            r.get("effortEstimate",""), r.get("complexity",""), r.get("decisionOwner",""), r.get("notes",""),
        ])
    for i in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(i)].width = 25
    ws.freeze_panes = "A2"

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    filename = f"Accenture_WRICEF_{date.today().strftime('%Y%m%d')}.xlsx"
    return StreamingResponse(buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'})


@router.post("/gap-analysis")
async def kdd_gap_analysis(request: Request, session_id: str = "default"):
    jp = kdd_data_path(session_id)
    if not jp.exists():
        raise HTTPException(404, "No KDD data. Generate KDDs first.")
    data = json.loads(jp.read_text("utf-8"))
    rows = data.get("rows", [])
    if not rows:
        raise HTTPException(400, "No KDD rows found.")

    async def gap_stream() -> AsyncGenerator[str, None]:
        # ── Step 1: compute stats from loaded rows ────────────────────────────
        yield sse({"type": "step", "step": 1})
        gap_rows  = [r for r in rows if r.get("fitGap") in ("Gap", "Partial Fit")]
        ricef_rows = [r for r in rows if r.get("ricefType") and r.get("ricefType") not in ("None", "")]
        open_rows  = [r for r in rows if (r.get("status") or "").lower() in ("open", "in progress", "tbc")]
        appr_rows  = [r for r in rows if (r.get("status") or "").lower() in ("approved", "closed", "complete")]
        red_rows   = [r for r in rows if r.get("ragStatus") == "Red"
                      or (r.get("complexity") == "High" and r.get("fitGap") in ("Gap", "Partial Fit"))]
        stats = {
            "totalRed":      len(red_rows),
            "totalGap":      len(gap_rows),
            "totalRICEF":    len(ricef_rows),
            "openDecisions": len(open_rows),
            "approved":      len(appr_rows),
        }
        yield sse({"type": "step_done", "step": 1})

        # ── Steps 2+3: single Claude call — risks + themes in one pass ──────────
        yield sse({"type": "step", "step": 2})
        gap_detail = "\n".join(
            f"[{r.get('kddId','')}] [{r.get('scopeItemId','')}] {r.get('scopeItemName','')} | "
            f"{r.get('designQuestion','')} | {r.get('fitGap','')} | complexity:{r.get('complexity','')} | module:{r.get('module','')}"
            for r in gap_rows[:35]
        )
        combined_prompt = f"""You are an SAP S/4HANA Cloud Public Edition expert. Analyse these KDD gaps in one pass.

Gap/Partial Fit KDDs ({len(gap_rows)} total, showing up to 35):
{gap_detail}

Return ONE valid JSON object — no markdown, no explanation:
{{
  "risks": [
    {{"severity":"High|Medium|Low","title":"...","scopeItems":["BD6","J59"],"integrationPoint":"...","description":"...","action":"..."}}
  ],
  "analysis": {{
    "healthSummary": "1-2 sentence overall health of the KDD log",
    "topAction": "single most important action before go-live",
    "workshopFocus": "what to focus on in workshops",
    "themes": [
      {{"priority":"Critical|High|Medium","theme":"...","items":[{{"kddId":"...","question":"...","approach":"specific Cloud PE Fiori app or BTP approach"}}]}}
    ]
  }}
}}

Rules: 3-5 risks. 2-4 themes with 2-5 items each. Cloud PE only — no ABAP/ECC/on-premise."""

        risks = []
        analysis = {}
        try:
            out = await spawn_claude(combined_prompt, timeout=150)
            s, e = out.find("{"), out.rfind("}")
            if s >= 0 and e > s:
                parsed = json.loads(out[s:e+1])
                risks    = parsed.get("risks", [])
                analysis = parsed.get("analysis", {})
            yield sse({"type": "step_done", "step": 2})
        except Exception as exc:
            yield sse({"type": "step_warn", "step": 2, "message": str(exc)[:100]})

        yield sse({"type": "step", "step": 3})
        yield sse({"type": "step_done", "step": 3})

        yield sse({"type": "complete", "stats": stats, "risks": risks, "analysis": analysis})

    return StreamingResponse(gap_stream(), media_type="text/event-stream",
                             headers={"Cache-Control": "no-cache", "X-Accel-Buffering": "no"})


@router.post("/critique")
async def kdd_critique(request: Request):
    body = await request.json()
    rows = body.get("rows", [])
    if not rows:
        raise HTTPException(400, "rows required")

    async def critique_stream() -> AsyncGenerator[str, None]:
        q_list = "\n".join(
            f"{i+1}. [{r.get('kddId','')}] [{r.get('fitGap','')}] {r.get('designQuestion','')}"
            for i, r in enumerate(rows)
        )
        prompt = f"""You are an SAP S/4HANA Cloud Public Edition expert reviewer.
Review these {len(rows)} KDD questions and identify:
1. Any on-premise or non-Cloud PE references (ABAP, ECC, SAP GUI, etc.)
2. Vague questions that need to name a specific Fiori app or BTP service
3. Missing coverage areas (Data Migration, Authorisation, Number Ranges, Org Structure)
4. Duplicate or overlapping questions

Questions:
{q_list}

Reply with a structured critique. Be concise. Cloud PE only."""
        yield sse({"type": "progress", "message": "Analysing KDD questions…"})
        try:
            out = await spawn_claude(prompt, timeout=90)
            yield sse({"type": "critique", "text": out.strip()})
        except Exception as exc:
            yield sse({"type": "error", "message": str(exc)})

    return StreamingResponse(critique_stream(), media_type="text/event-stream",
                             headers={"Cache-Control": "no-cache", "X-Accel-Buffering": "no"})
