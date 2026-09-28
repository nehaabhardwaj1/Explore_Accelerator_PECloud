# bdcq_agent/routes.py — FastAPI router for all /api/bdcq/* endpoints
import asyncio
import io
import json
import logging
import os
import re
import sys
import tempfile
from pathlib import Path
from typing import Optional

from fastapi import APIRouter, File, HTTPException, Request, UploadFile
from fastapi.responses import JSONResponse, StreamingResponse

ROOT   = Path(__file__).resolve().parent.parent
log    = logging.getLogger(__name__)
router = APIRouter(prefix="/api/bdcq", tags=["bdcq"])

BDCQ_PATH          = ROOT / "bdcq" / "bdcq-questions.json"
GROUNDING_DIR      = ROOT / "bdcq" / "grounding"
ALL_COUNTRIES_PATH = GROUNDING_DIR / "BDCQ_S4HANA_PublicCloud_v4_AllCountries.json"


def _slug(name: str) -> str:
    return re.sub(r"[^\w-]", "_", Path(name).stem)


def _grounding_files() -> list:
    if not GROUNDING_DIR.exists():
        return []
    files = []
    for p in sorted(GROUNDING_DIR.glob("*.json")):
        try:
            data = json.loads(p.read_text("utf-8"))
            files.append({**data, "slug": _slug(p.name), "filename": p.name})
        except Exception:
            pass
    return files


# ── File I/O helpers ──────────────────────────────────────────────────────────

def _read_bdcq():
    if not BDCQ_PATH.exists():
        return None
    return json.loads(BDCQ_PATH.read_text("utf-8"))



def _read_all_countries():
    if not ALL_COUNTRIES_PATH.exists():
        return None
    data = json.loads(ALL_COUNTRIES_PATH.read_text("utf-8"))
    return data if isinstance(data, list) else data.get("rows") or data.get("data")


# ── Country / Industry normalisation (imported from localization module) ───────

from bdcq_enricher.localization import (
    COUNTRY_RULES, INDUSTRY_RULES, CLOUDPE_CONSTRAINTS,
    normalize_countries, normalize_industry,
)

_ISO_DISPLAY = {
    "IN": "India", "DE": "Germany", "US": "USA", "GB": "UK", "BR": "Brazil",
    "FR": "France", "AU": "Australia", "SG": "Singapore", "JP": "Japan",
    "CN": "China", "KR": "South Korea", "VN": "Vietnam", "AE": "UAE",
    "SA": "Saudi Arabia", "NL": "Netherlands", "PL": "Poland", "IT": "Italy",
    "ES": "Spain", "AT": "Austria", "SE": "Sweden", "MX": "Mexico",
    "CA": "Canada", "ZA": "South Africa", "AR": "Argentina",
}


# ── Claude subprocess (same pattern as generation_agent.py) ───────────────────

def _resolve_claude() -> str:
    import os
    if sys.platform != "win32":
        return "claude"
    if (p := os.environ.get("CLAUDE_PATH")) and Path(p).exists():
        return p
    home = os.environ.get("USERPROFILE") or str(Path.home())
    appdata = os.environ.get("APPDATA") or str(Path(home) / "AppData" / "Roaming")
    candidates = [
        Path(home) / ".local" / "bin" / "claude.exe",
        Path(home) / "AppData" / "Local" / "Programs" / "claude" / "claude.exe",
        Path(appdata) / "npm" / "claude.exe",
        Path(appdata) / "npm" / "claude.cmd",
    ]
    for p in candidates:
        if p.exists():
            return str(p)
    return "claude"


_CLAUDE_EXE = _resolve_claude()
CLAUDE_MODEL = os.environ.get("CLAUDE_MODEL", "claude-opus-4-5")


def _make_env() -> dict:
    import os
    home = str(Path.home())
    appdata = os.environ.get("APPDATA") or str(Path(home) / "AppData" / "Roaming")
    sep  = ";" if sys.platform == "win32" else ":"
    extra = sep.join([
        str(Path(home) / ".local" / "bin"),
        str(Path(home) / "AppData" / "Local" / "Programs" / "claude"),
        str(Path(appdata) / "npm"),
        os.environ.get("PATH", ""),
    ])
    return {**dict(os.environ), "PATH": extra}


async def _spawn_claude(prompt: str, timeout: int = 120) -> str:
    is_cmd = sys.platform == "win32" and _CLAUDE_EXE.lower().endswith((".cmd", ".bat"))
    if is_cmd:
        proc = await asyncio.create_subprocess_shell(
            f'"{_CLAUDE_EXE}" --model {CLAUDE_MODEL} --print --dangerously-skip-permissions',
            stdin=asyncio.subprocess.PIPE, stdout=asyncio.subprocess.PIPE,
            stderr=asyncio.subprocess.PIPE, env=_make_env(), cwd=tempfile.gettempdir(),
        )
    else:
        proc = await asyncio.create_subprocess_exec(
            _CLAUDE_EXE, "--model", CLAUDE_MODEL, "--print", "--dangerously-skip-permissions",
            stdin=asyncio.subprocess.PIPE, stdout=asyncio.subprocess.PIPE,
            stderr=asyncio.subprocess.PIPE, env=_make_env(), cwd=tempfile.gettempdir(),
        )
    try:
        stdout, _ = await asyncio.wait_for(proc.communicate(prompt.encode("utf-8")), timeout=timeout)
    except asyncio.TimeoutError:
        proc.kill()
        raise TimeoutError(f"Claude timed out after {timeout}s")
    if proc.returncode != 0:
        raise RuntimeError(f"Claude exited {proc.returncode}")
    return stdout.decode("utf-8", errors="replace")


# ── GET /api/bdcq/status ──────────────────────────────────────────────────────

@router.get("/status")
async def bdcq_status():
    data = _read_bdcq()
    if not data:
        return {"found": False, "loaded": False, "totalDomains": 0, "totalQuestions": 0, "domains": []}
    domains = data.get("domains", [])
    total   = sum(sum(len(s.get("rows", [])) for s in d.get("sheets", [])) for d in domains)

    # Build domainMeta: per-domain sheet summary for the domain selector UI
    domain_meta = []
    all_countries: set = set()
    for d in domains:
        sheets = d.get("sheets", [])
        sheet_names = [s.get("sheet", "") for s in sheets]
        for sn in sheet_names:
            snu = sn.upper()
            if len(snu) == 2 and snu.isalpha():
                all_countries.add(sn)
        # Collect unique column headers across all sheets in this domain
        col_set: list = []
        col_seen: set = set()
        for s in sheets:
            for h in s.get("headers", []) or []:
                if h and h not in col_seen:
                    col_seen.add(h)
                    col_set.append(h)
            # Also derive from first row keys if headers absent
            if not col_set:
                for row in s.get("rows", [])[:1]:
                    for k in row.keys():
                        if k and k not in col_seen:
                            col_seen.add(k)
                            col_set.append(k)
        domain_meta.append({
            "domain":   d.get("domain", ""),
            "sheets":   sheet_names,
            "rowCount": sum(len(s.get("rows", [])) for s in sheets),
            "columns":  col_set,
        })

    return {
        "found":          True,
        "loaded":         True,
        "totalDomains":   len(domains),
        "totalQuestions": total,
        "generatedAt":    data.get("generatedAt"),
        "domains":        [d.get("domain") for d in domains],
        "domainMeta":     domain_meta,
        "countries":      sorted(all_countries),
    }


# ── POST /api/bdcq/upload (JSON body from Chrome extension) ──────────────────

@router.post("/upload")
async def bdcq_upload(request: Request):
    body = await request.json()
    BDCQ_PATH.parent.mkdir(exist_ok=True)
    BDCQ_PATH.write_text(json.dumps(body, indent=2, ensure_ascii=False), "utf-8")
    domains = body.get("domains", [])
    total   = sum(sum(len(s.get("rows", [])) for s in d.get("sheets", [])) for d in domains)
    return {"ok": True, "domainCount": len(domains), "totalQuestions": total}


# ── GET /api/bdcq/grounding ────────────────────────────────────────────────────

@router.get("/grounding")
async def get_grounding():
    files = _grounding_files()
    if not files:
        return {"exists": False, "files": [], "totalRows": 0}
    total = sum(f.get("totalRows", 0) for f in files)
    return {"exists": True, "files": files, "totalRows": total}


# ── POST /api/bdcq/grounding/upload ──────────────────────────────────────────

@router.post("/grounding/upload")
async def upload_grounding(file: UploadFile = File(...)):
    from datetime import datetime
    import openpyxl
    content   = await file.read()
    wb        = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    sheets    = []
    for sheet_name in wb.sheetnames:
        ws        = wb[sheet_name]
        rows_iter = list(ws.iter_rows(values_only=True))
        if not rows_iter:
            continue
        headers    = [str(h or "").strip() for h in rows_iter[0]]
        sheet_rows = []
        for row in rows_iter[1:]:
            obj = {headers[i]: (str(v) if v is not None else "") for i, v in enumerate(row) if i < len(headers)}
            if any(v for v in obj.values()):
                sheet_rows.append(obj)
        if sheet_rows:
            sheets.append({
                "sheet": sheet_name, "source": file.filename,
                "columns": headers, "rowCount": len(sheet_rows), "rows": sheet_rows,
            })
    slug     = _slug(file.filename or "upload")
    out_path = GROUNDING_DIR / f"{slug}.json"
    GROUNDING_DIR.mkdir(parents=True, exist_ok=True)
    payload  = {
        "origName":   file.filename,
        "uploadedAt": datetime.utcnow().isoformat() + "Z",
        "totalRows":  sum(len(s["rows"]) for s in sheets),
        "sheets":     sheets,
    }
    out_path.write_text(json.dumps(payload, indent=2, ensure_ascii=False), "utf-8")
    return {"ok": True, "slug": slug, "sheetCount": len(sheets),
            "totalRows": payload["totalRows"]}


# ── DELETE /api/bdcq/grounding/{slug} ────────────────────────────────────────

@router.delete("/grounding/{slug}")
async def delete_grounding(slug: str):
    safe = re.sub(r"[^\w-]", "_", slug)
    target = GROUNDING_DIR / f"{safe}.json"
    if target.exists():
        target.unlink()
    return {"ok": True}


# ── POST /api/bdcq/preview ────────────────────────────────────────────────────

@router.post("/preview")
async def bdcq_preview(request: Request):
    body = await request.json()
    scope_ids        = body.get("scopeIds", [])
    selected_domains = body.get("selectedDomains", [])
    selected_columns = body.get("selectedColumns", [])
    countries        = body.get("countries", [])

    data = _read_bdcq()
    if not data:
        raise HTTPException(404, "bdcq-questions.json not found. Upload the file first.")

    scope_set   = {s.upper().strip() for s in scope_ids if s}
    country_set = {c.upper().strip() for c in countries if c}
    domain_set  = {d.lower().strip() for d in selected_domains if d}

    def is_country_sheet(name):
        return bool(re.match(r"^[A-Za-z]{2}$", name)) or name == "LC"

    all_rows    = []
    all_headers = []
    header_set  = set()

    for domain in data.get("domains", []):
        domain_name = domain.get("domain", "")
        if domain_set and domain_name.lower().strip() not in domain_set:
            continue
        for sheet in domain.get("sheets", []):
            sheet_name = sheet.get("sheet", "")
            if re.match(r"template|change history|overview|instructions|cover|readme", sheet_name, re.IGNORECASE):
                continue
            if is_country_sheet(sheet_name):
                if country_set and sheet_name.upper() not in country_set:
                    continue
            for h in sheet.get("headers", []):
                if h and h not in header_set:
                    header_set.add(h)
                    all_headers.append(h)
            for row in sheet.get("rows", []):
                if scope_set:
                    vals = {str(v or "").upper().strip() for v in row.values()}
                    if not vals.intersection(scope_set):
                        continue
                all_rows.append({
                    "Module":  domain_name,
                    "Section": sheet_name,
                    "_localizationCountry": sheet_name.upper() if is_country_sheet(sheet_name) else "Global",
                    **row,
                })

    # Inject AllCountries grounding rows
    if country_set:
        ac_rows = _read_all_countries()
        if ac_rows:
            _COUNTRY_NORM = {
                "india":"IN","germany":"DE","usa":"US","uk":"GB","brazil":"BR",
                "france":"FR","australia":"AU","singapore":"SG","japan":"JP",
                "china":"CN","south korea":"KR","vietnam":"VN","netherlands":"NL",
                "uae":"AE","saudi arabia":"SA",
            }
            seen_q = set()
            for row in ac_rows:
                cr      = (row.get("Country / Region") or "").strip()
                question= (row.get("Question") or "").strip()
                if not cr or cr == "Global" or not question:
                    continue
                country_part      = re.sub(r"^Local\s*[–\-]\s*", "", cr, flags=re.IGNORECASE).strip()
                is_all_countries  = country_part.lower() == "all countries"
                section_code      = None
                if is_all_countries:
                    dk = f"LC::{question}"
                    if dk in seen_q: continue
                    seen_q.add(dk)
                    section_code = "LC"
                else:
                    for part in re.split(r"[,;]", country_part):
                        part = part.strip().lower()
                        iso  = _COUNTRY_NORM.get(part) or part.upper()[:2]
                        if iso in country_set:
                            dk = f"{iso}::{question}"
                            if dk not in seen_q:
                                seen_q.add(dk)
                                section_code = iso
                            break
                if not section_code:
                    continue
                bpa     = (row.get("Business Process Area") or "").lower()
                matches = not domain_set
                if not matches:
                    for d in domain_set:
                        if d in bpa or bpa in d:
                            matches = True; break
                if not matches:
                    continue
                all_rows.append({
                    "Module":               row.get("Business Process Area", "General"),
                    "Section":              section_code,
                    "_localizationCountry": section_code,
                    "_fromAllCountries":    True,
                    "Question":             question,
                    "Topic":                row.get("Topic", ""),
                    "Country / Region":     "Local – All Countries" if section_code == "LC"
                                            else f"Local – {_ISO_DISPLAY.get(section_code, section_code)}",
                })
                for h in ["Question","Topic","Topic Definition","SAP Workstream","Localization Type","Sample Value"]:
                    if h not in header_set:
                        header_set.add(h)
                        all_headers.append(h)

    base_cols = ["Module", "Section"]
    data_cols = [all_headers[i] for i in range(len(all_headers))
                 if selected_columns and all_headers[i].lower() in {c.lower() for c in selected_columns}] \
                if selected_columns else all_headers
    output_cols = base_cols + data_cols

    question_col = next((c for c in data_cols if "question" in c.lower()), data_cols[0] if data_cols else "Question")

    return {"ok": True, "rows": all_rows, "columns": output_cols, "questionCol": question_col, "rowCount": len(all_rows)}


# ── POST /api/bdcq/enrich ─────────────────────────────────────────────────────

@router.post("/enrich")
async def bdcq_enrich(request: Request):
    body        = await request.json()
    domain      = body.get("domain", "")
    question_col= body.get("questionCol", "Question")
    raw_rows    = body.get("rows", [])
    raw_country = body.get("country")
    raw_industry= body.get("industry", "")

    if not raw_rows:
        return {"ok": True, "enrichments": []}

    def get_col(row, *candidates):
        for c in candidates:
            for k, v in row.items():
                if c.lower() in k.lower() and str(v or "").strip():
                    return str(v).strip()
        return ""

    domain_lower = (domain or "").strip().lower()
    def is_invalid(r):
        q = str(r.get(question_col) or r.get("Question") or "").strip()
        return not q or len(q) < 3 or q.lower() == domain_lower

    rows         = [r for r in raw_rows if not is_invalid(r)]
    valid_indices= [not is_invalid(r) for r in raw_rows]

    if not rows:
        return {"ok": True, "enrichments": [{"standardValue": "", "watchOut": ""} for _ in raw_rows]}

    # Build question lines
    question_lines = []
    for i, r in enumerate(rows):
        question = str(r.get(question_col) or r.get("Question") or "").strip()[:300]
        topic_def= get_col(r, "topic definition", "definition", "description")[:400]
        std_val  = get_col(r, "solution", "sample value", "standard value", "standard values")[:600]
        watch_out= get_col(r, "watch out", "constraint", "watch-out")[:500]
        sscui    = get_col(r, "sscui reference", "sscui", "configuration activity")[:150]
        sap_id   = get_col(r, "sap id", "sapid", "id")[:20]
        topic    = get_col(r, "topic")[:100]
        block    = f"{i+1}. Question: {question}"
        if topic:     block += f"\n   Topic: {topic}"
        if sscui:     block += f"\n   SSCUI: {sscui}"
        if sap_id:    block += f"\n   SAP ID: {sap_id}"
        if topic_def: block += f"\n   Context: {topic_def}"
        if std_val:   block += f"\n   SAP Standard Value (from BDCQ): {std_val}"
        if watch_out: block += f"\n   SAP Watch-Out (from BDCQ): {watch_out}"
        question_lines.append(block)
    q_block = "\n\n".join(question_lines)

    has_ref = any(get_col(r, "solution", "sample value", "standard value") or get_col(r, "watch out", "constraint") for r in rows)

    # Country context
    country_codes = []
    if isinstance(raw_country, list):
        for c in raw_country:
            country_codes.extend(normalize_countries(c))
    else:
        country_codes.extend(normalize_countries(raw_country or ""))

    rows_text = " ".join(
        " ".join([get_col(r, "process area"), get_col(r, "sscui reference", "sscui"), get_col(r, "topic"), str(r.get(question_col) or "")]).lower()
        for r in rows
    )

    country_block = ""
    if country_codes:
        lines = []
        for code in country_codes:
            cd = COUNTRY_RULES.get(code)
            if not cd: continue
            matching = [rule for rule in cd["rules"] if any(kw.lower() in rows_text for kw in rule["processAreaKeywords"])]
            if not matching: continue
            lines.append(f"\n{cd['name']}:")
            for rule in matching:
                lines.append(f"  • [{'MANDATORY' if rule['mandatory'] else 'Conditional'}] {rule['requirement']}")
                lines.append(f"    Basis: {rule['basis']}")
                lines.append(f"    Cloud PE approach: {rule['cloudPEApproach']}")
                lines.append(f"    Confirm with client: {rule['confirmPrompt']}")
        if lines:
            country_block = "\n\nCOUNTRY LOCALISATION CONTEXT (" + ", ".join(country_codes) + "):\n" + "\n".join(lines)

    # Industry context
    industry_key  = normalize_industry(raw_industry)
    industry_block= ""
    if industry_key and industry_key in INDUSTRY_RULES:
        ind    = INDUSTRY_RULES[industry_key]
        themes = [f"  • {v}" for k, v in ind["processAreaThemes"].items()
                  if k == "default" or k in domain.lower() or k in rows_text]
        industry_block = (
            f"\n\nINDUSTRY CONTEXT — {ind['name']}:\n"
            f"Frame ALL answers from a {ind['name']} industry perspective.\n"
            + ("\n".join(themes) if themes else f"  • {ind['processAreaThemes']['default']}")
            + "\nIndustry-specific questions to raise where applicable:\n"
            + "\n".join(f"  → {q}" for q in ind["mustAsk"])
        )

    # Cloud PE constraints
    relevant = [c for c in CLOUDPE_CONSTRAINTS if any(kw in rows_text for kw in c["keywords"])]
    cloud_block = ""
    if relevant:
        cloud_block = "\n\nCLOUD PE CONSTRAINT WATCH-OUTS for this batch:\n" + "\n".join(
            f'  [{c["type"]}] Client may expect: "{c["expectation"]}" → Reality: {c["reality"]} Extension: {c["extensionPath"]}'
            for c in relevant
        )

    # Grounding
    grounding_block = ""
    all_grounding_sheets = [s for f in _grounding_files() for s in f.get("sheets", [])]
    if all_grounding_sheets:
        words  = list(set(re.findall(r"\b\w{3,}\b", rows_text.lower())))
        scored = []
        for sheet in all_grounding_sheets:
            bonus = 3 if any(w in sheet["sheet"].lower() for w in words if len(w) >= 3) else 0
            for row in sheet.get("rows", []):
                row_str = " ".join(str(v) for v in row.values()).lower()
                hits    = sum(1 for w in words if w in row_str) + bonus
                if hits > 0:
                    scored.append((hits, sheet["sheet"], row))
        scored.sort(key=lambda x: -x[0])
        top = scored[:10]
        if top:
            lines = [f"  [{s}] " + " | ".join(f"{k}: {str(v)[:300]}" for k, v in r.items() if v)
                     for _, s, r in top]
            grounding_block = (
                "\n\nREFERENCE GROUNDING DATA (from client's uploaded reference Excel):\n"
                "Use the rows below as the PRIMARY source for standardValue where they directly answer a question.\n"
                + "\n".join(lines)
            )

    ref_instruction = (
        """Where SAP BDCQ reference data is provided, reformat it:
- "standardValue": 1–3 concise sentences. Incorporate country statutory requirements where applicable. Add industry-specific framing. Keep any specific codes/types. Name the Fiori app.
- "watchOut": 1–3 bullet points. Include country statutory deadlines or compliance risks. Flag Cloud PE constraints where relevant.
Where reference data is absent for a question, use your SAP Cloud PE knowledge."""
        if has_ref else
        f"""For each question:
- "standardValue": Best available SAP Cloud PE standard or default. Incorporate country statutory requirements for MANDATORY items as pre-answered context. Frame from {industry_key or "the client"} industry perspective. Include specific codes/types where known. Name the Fiori app. If genuinely client-specific, say "Client-specific — confirm in workshop".
- "watchOut": Key constraints or irreversible decisions. Include country statutory deadlines or compliance risks. Flag Cloud PE constraints. Use "" if none."""
    )

    prompt = f"""You are an SAP S/4HANA Cloud Public Edition (Cloud PE) configuration consultant.

For each question, use the Topic, SSCUI reference, and SAP ID provided to give the most specific answer you can from your SAP knowledge.{grounding_block}{country_block}{industry_block}{cloud_block}

{ref_instruction}

RULES — strictly enforced:
- SAP S/4HANA Cloud Public Edition ONLY. No T-codes, SPRO, ABAP, SAP GUI, or on-premise references.
- Name the specific Fiori app (e.g. "Manage G/L Accounts", "Create Supplier Invoice") not just "SAP Fiori".
- For MANDATORY country requirements: pre-answer in standardValue, do not leave as open question.

Questions to enrich:

{q_block}

Return ONLY a JSON array of exactly {len(rows)} objects:
[{{"standardValue": "...", "watchOut": "..."}}]
JSON only. No markdown."""

    try:
        out  = await _spawn_claude(prompt, timeout=120)
        s    = out.find("[")
        e    = out.rfind("]")
        if s < 0 or e < 0:
            raise ValueError("No JSON array in response")
        parsed = json.loads(out[s:e+1])
        if not isinstance(parsed, list) or len(parsed) != len(rows):
            raise ValueError(f"Expected {len(rows)} enrichments, got {len(parsed) if isinstance(parsed, list) else type(parsed)}")
    except Exception as exc:
        log.warning("BDCQ enrich Claude call failed: %s", exc)
        return {"ok": False, "error": str(exc), "enrichments": [{"standardValue": "", "watchOut": ""} for _ in raw_rows]}

    # Map back to raw_rows indices
    enrichments = []
    parsed_idx  = 0
    for is_valid in valid_indices:
        if is_valid and parsed_idx < len(parsed):
            enrichments.append(parsed[parsed_idx])
            parsed_idx += 1
        else:
            enrichments.append({"standardValue": "", "watchOut": ""})

    return {"ok": True, "enrichments": enrichments}


# ── POST /api/bdcq/ai-local-questions ────────────────────────────────────────

@router.post("/ai-local-questions")
async def bdcq_ai_local_questions(request: Request):
    body     = await request.json()
    domain   = body.get("domain", "")
    countries= body.get("countries", [])
    industry = body.get("industry", "")

    if not domain:
        raise HTTPException(400, "domain is required")

    country_list = ", ".join(countries) if countries else "India"
    industry_str = f" for a {industry} company" if industry else ""

    prompt = f"""You are an SAP S/4HANA Cloud Public Edition configuration consultant specialising in localisation.

Generate 8–12 country-specific BDCQ configuration questions for the domain "{domain}"{industry_str}, covering statutory requirements for: {country_list}.

For each question return:
- "Question": the configuration question (specific, actionable)
- "Country / Region": ISO country code(s) or "All Countries"
- "Topic": short topic label (5–8 words)
- "Sample Value": the typical SAP Cloud PE standard value or approach
- "Level": always "L3"
- "_aiGenerated": true
- "_localizationCountry": same as Country / Region

Focus on: tax (GST/VAT/withholding), statutory reports, e-invoicing mandates, payroll compliance, bank formats, data residency — whichever apply to {domain}.
Only include questions where the configuration decision differs by country.
Use SAP S/4HANA Cloud Public Edition terminology — no T-codes, no ABAP, no on-premise references.

Return ONLY a JSON array. No markdown, no explanation.
[{{"Question":"...","Country / Region":"...","Topic":"...","Sample Value":"...","Level":"L3","_aiGenerated":true,"_localizationCountry":"..."}}]"""

    try:
        out = await _spawn_claude(prompt, timeout=90)
        s = out.find("[")
        e = out.rfind("]")
        if s < 0 or e < 0:
            raise ValueError("No JSON array in Claude response")
        rows = json.loads(out[s:e+1])
        if not isinstance(rows, list):
            raise ValueError("Expected JSON array")
    except Exception as exc:
        log.warning("ai-local-questions Claude call failed: %s", exc)
        return JSONResponse({"ok": False, "error": str(exc)}, status_code=200)

    return {"ok": True, "rows": rows, "domain": domain}


# ── POST /api/bdcq/upload-grounding ──────────────────────────────────────────
# Accepts { files: [{ filename, content: base64 }] }, parses each Excel file,
# saves to bdcq/grounding/<slug>.json, returns { ok, saved: [...] }.

@router.post("/upload-grounding")
async def upload_grounding_b64(request: Request):
    import base64
    import openpyxl
    from datetime import datetime

    body = await request.json()
    file_list = body.get("files", [])
    if not file_list:
        raise HTTPException(400, "No files provided")

    GROUNDING_DIR.mkdir(parents=True, exist_ok=True)
    saved = []

    for entry in file_list:
        filename = entry.get("filename", "upload.xlsx")
        b64_content = entry.get("content", "")
        if not b64_content:
            continue
        try:
            raw = base64.b64decode(b64_content)
            wb  = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        except Exception as exc:
            saved.append({"filename": filename, "error": str(exc)})
            continue

        sheets = []
        for sheet_name in wb.sheetnames:
            ws        = wb[sheet_name]
            rows_iter = list(ws.iter_rows(values_only=True))
            if not rows_iter:
                continue
            headers    = [str(h or "").strip() for h in rows_iter[0]]
            sheet_rows = []
            for row in rows_iter[1:]:
                obj = {headers[i]: (str(v) if v is not None else "")
                       for i, v in enumerate(row) if i < len(headers)}
                if any(v for v in obj.values()):
                    sheet_rows.append(obj)
            if sheet_rows:
                sheets.append({
                    "sheet": sheet_name, "source": filename,
                    "columns": headers, "rowCount": len(sheet_rows), "rows": sheet_rows,
                })

        slug     = _slug(filename)
        out_path = GROUNDING_DIR / f"{slug}.json"
        total_rows = sum(len(s["rows"]) for s in sheets)
        payload = {
            "origName":   filename,
            "uploadedAt": datetime.utcnow().isoformat() + "Z",
            "totalRows":  total_rows,
            "sheets":     sheets,
        }
        out_path.write_text(json.dumps(payload, indent=2, ensure_ascii=False), "utf-8")
        saved.append({"filename": filename, "slug": slug,
                      "sheetCount": len(sheets), "totalRows": total_rows})

    return {"ok": True, "saved": saved}


# ── POST /api/bdcq/export-editor ─────────────────────────────────────────────
# Called by bdcqExportExcel() in the editor tab.
# Saves the file to output/ and returns JSON {ok, filename, rowCount} so the
# frontend can offer a download link without a streaming response.

@router.post("/export-editor")
async def bdcq_export_editor(request: Request):
    import base64
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    from datetime import date

    body             = await request.json()
    rows             = body.get("rows", [])
    columns          = body.get("columns", [])
    client_name      = body.get("clientName", "")
    overview_text    = body.get("overviewText", "")
    filename_stem    = body.get("filename", "")
    client_logo_b64  = body.get("clientLogo")
    partner_logo_b64 = body.get("partnerLogo")

    if not rows:
        raise HTTPException(400, "No rows to export")

    today = date.today().strftime("%Y%m%d")
    if not filename_stem:
        safe = re.sub(r"[^\w]", "_", client_name or "Client")[:20]
        filename_stem = f"{safe}_BDCQ_{today}"
    filename = re.sub(r"[^\w.\-]", "_", filename_stem) + ".xlsx"

    wb = openpyxl.Workbook()

    # ── Cover sheet ──────────────────────────────────────────────────────────
    cover = wb.active
    cover.title = "Cover"
    cover.sheet_view.showGridLines = False
    cover.column_dimensions["A"].width = 80

    cover_fill = PatternFill("solid", fgColor="1F4E79")
    title_font = Font(name="Calibri", color="FFFFFF", bold=True, size=20)
    sub_font   = Font(name="Calibri", color="DDEEFF", size=13)
    body_font  = Font(name="Calibri", color="1F2937", size=11)

    cover["A1"] = "Business Driven Configuration Questionnaire (BDCQ)"
    cover["A1"].font = title_font
    cover["A1"].fill = cover_fill
    cover["A1"].alignment = Alignment(wrap_text=True, vertical="center", indent=2)
    cover.row_dimensions[1].height = 60

    cover["A2"] = f"Client: {client_name}" if client_name else "SAP S/4HANA Cloud — Public Edition"
    cover["A2"].font = sub_font
    cover["A2"].fill = PatternFill("solid", fgColor="2563EB")
    cover["A2"].alignment = Alignment(vertical="center", indent=2)
    cover.row_dimensions[2].height = 32

    if overview_text:
        cover["A3"] = overview_text
        cover["A3"].font = body_font
        cover["A3"].alignment = Alignment(wrap_text=True, vertical="top", indent=2)
        cover.row_dimensions[3].height = max(80, len(overview_text) // 3)

    # Embed logos if provided (base64 PNG/JPEG)
    try:
        from openpyxl.drawing.image import Image as XLImage
        img_row = 4 if overview_text else 3
        col_offset = 0
        for b64_data, label in [(client_logo_b64, "client"), (partner_logo_b64, "partner")]:
            if not b64_data:
                continue
            raw = b64_data.split(",", 1)[-1]  # strip data:image/...;base64, prefix
            img_bytes = base64.b64decode(raw)
            img_buf = io.BytesIO(img_bytes)
            xl_img = XLImage(img_buf)
            xl_img.width  = 160
            xl_img.height = 80
            anchor = f"{'A' if col_offset == 0 else 'C'}{img_row}"
            cover.add_image(xl_img, anchor)
            col_offset += 1
    except Exception:
        pass  # logo embedding is optional — skip silently if Pillow not installed

    # ── Data sheet ───────────────────────────────────────────────────────────
    ws = wb.create_sheet("BDCQ Questions")
    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(color="FFFFFF", bold=True, size=9)

    ws.append(columns)
    for cell in ws[1]:
        cell.font  = header_font
        cell.fill  = header_fill
        cell.alignment = Alignment(wrap_text=True, vertical="center")

    for row in rows:
        ws.append([str(row.get(col) or "") for col in columns])
        for cell in ws[ws.max_row]:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    for i in range(1, len(columns) + 1):
        ws.column_dimensions[get_column_letter(i)].width = 38
    ws.freeze_panes = "A2"

    # ── Save to output/ ───────────────────────────────────────────────────────
    from routes.deps import OUTPUT_DIR
    out_path = OUTPUT_DIR / filename
    wb.save(str(out_path))

    return {"ok": True, "filename": filename, "rowCount": len(rows)}


# ── POST /api/bdcq/export ─────────────────────────────────────────────────────

@router.post("/export")
async def bdcq_export(request: Request):
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    body    = await request.json()
    rows    = body.get("rows", [])
    columns = body.get("columns", [])
    client  = body.get("client", "")
    domain  = body.get("domain", "")

    if not rows:
        raise HTTPException(400, "No rows to export")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "BDCQ"

    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(color="FFFFFF", bold=True, size=9)

    ws.append(columns)
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(wrap_text=True, vertical="center")

    for row in rows:
        ws.append([str(row.get(col) or "") for col in columns])
        for cell in ws[ws.max_row]:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    for i in range(1, len(columns) + 1):
        ws.column_dimensions[get_column_letter(i)].width = 40

    ws.freeze_panes = "A2"

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    from datetime import date
    safe_client = re.sub(r"[^\w]", "_", client or "")[:15]
    safe_domain = re.sub(r"[^\w]", "_", domain or "")[:15]
    today = date.today().strftime("%Y%m%d")
    filename = f"BDCQ{'_'+safe_client if safe_client else ''}{'_'+safe_domain if safe_domain else ''}_{today}.xlsx"

    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


# ── POST /api/bdcq/generate ───────────────────────────────────────────────────
# Quick Export: filter rows (same as /preview) + build Excel + save to output/.
# Returns JSON {ok, filename, rowCount} so the frontend can offer a download link.

@router.post("/generate")
async def bdcq_generate(request: Request):
    import base64
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    from datetime import date
    from routes.deps import OUTPUT_DIR

    body             = await request.json()
    filename_stem    = (body.get("filename") or "").strip()
    client_name      = (body.get("clientName") or "").strip()
    scope_ids        = body.get("scopeIds", [])
    selected_domains = body.get("selectedDomains", [])
    selected_columns = body.get("selectedColumns", [])
    countries        = body.get("countries", [])
    overview_text    = (body.get("overviewText") or "").strip()
    client_logo_b64  = body.get("clientLogo")
    partner_logo_b64 = body.get("partnerLogo")

    bdcq_data = _read_bdcq()
    if not bdcq_data:
        raise HTTPException(404, "bdcq-questions.json not found. Upload BDCQ files first.")

    # ── Build rows (same logic as /preview) ──────────────────────────────────
    scope_set   = {s.upper().strip() for s in scope_ids if s}
    country_set = {c.upper().strip() for c in countries if c}
    domain_set  = {d.lower().strip() for d in selected_domains if d}

    def is_country_sheet(name):
        return bool(re.match(r"^[A-Za-z]{2}$", name)) or name == "LC"

    all_rows, all_headers, header_set = [], [], set()

    for domain in bdcq_data.get("domains", []):
        domain_name = domain.get("domain", "")
        if domain_set and domain_name.lower().strip() not in domain_set:
            continue
        for sheet in domain.get("sheets", []):
            sheet_name = sheet.get("sheet", "")
            if re.match(r"template|change history|overview|instructions|cover|readme", sheet_name, re.IGNORECASE):
                continue
            if is_country_sheet(sheet_name):
                if country_set and sheet_name.upper() not in country_set:
                    continue
            for h in sheet.get("headers", []):
                if h and h not in header_set:
                    header_set.add(h)
                    all_headers.append(h)
            for row in sheet.get("rows", []):
                if scope_set:
                    vals = {str(v or "").upper().strip() for v in row.values()}
                    if not vals.intersection(scope_set):
                        continue
                all_rows.append({
                    "Module":  domain_name,
                    "Section": sheet_name,
                    "_localizationCountry": sheet_name.upper() if is_country_sheet(sheet_name) else "Global",
                    **row,
                })

    if country_set:
        ac_rows = _read_all_countries()
        if ac_rows:
            _COUNTRY_NORM = {
                "india":"IN","germany":"DE","usa":"US","uk":"GB","brazil":"BR",
                "france":"FR","australia":"AU","singapore":"SG","japan":"JP",
                "china":"CN","south korea":"KR","vietnam":"VN","netherlands":"NL",
                "uae":"AE","saudi arabia":"SA",
            }
            seen_q: set = set()
            for row in ac_rows:
                cr       = (row.get("Country / Region") or "").strip()
                question = (row.get("Question") or "").strip()
                if not cr or cr == "Global" or not question:
                    continue
                country_part     = re.sub(r"^Local\s*[–\-]\s*", "", cr, flags=re.IGNORECASE).strip()
                is_all_countries = country_part.lower() == "all countries"
                section_code     = None
                if is_all_countries:
                    dk = f"LC::{question}"
                    if dk in seen_q: continue
                    seen_q.add(dk); section_code = "LC"
                else:
                    for part in re.split(r"[,;]", country_part):
                        part = part.strip().lower()
                        iso  = _COUNTRY_NORM.get(part) or part.upper()[:2]
                        if iso in country_set:
                            dk = f"{iso}::{question}"
                            if dk not in seen_q:
                                seen_q.add(dk); section_code = iso
                            break
                if not section_code:
                    continue
                bpa = (row.get("Business Process Area") or "").lower()
                matches = not domain_set or any(d in bpa or bpa in d for d in domain_set)
                if not matches:
                    continue
                all_rows.append({
                    "Module": row.get("Business Process Area", "General"),
                    "Section": section_code,
                    "_localizationCountry": section_code,
                    "_fromAllCountries": True,
                    "Question": question,
                    "Topic": row.get("Topic", ""),
                    "Country / Region": "Local – All Countries" if section_code == "LC"
                                        else f"Local – {_ISO_DISPLAY.get(section_code, section_code)}",
                })
                for h in ["Question","Topic","Topic Definition","SAP Workstream","Localization Type","Sample Value"]:
                    if h not in header_set:
                        header_set.add(h); all_headers.append(h)

    if not all_rows:
        raise HTTPException(400, "No BDCQ rows matched the selected domains/countries.")

    base_cols = ["Module", "Section"]
    data_cols = [h for h in all_headers
                 if selected_columns and h.lower() in {c.lower() for c in selected_columns}] \
                if selected_columns else all_headers
    columns = base_cols + data_cols

    # ── Build Excel ───────────────────────────────────────────────────────────
    wb = openpyxl.Workbook()

    # Cover sheet
    cover = wb.active
    cover.title = "Cover"
    cover.sheet_view.showGridLines = False
    cover.column_dimensions["A"].width = 80
    cover_fill = PatternFill("solid", fgColor="1F4E79")
    cover["A1"] = "Business Driven Configuration Questionnaire (BDCQ)"
    cover["A1"].font = Font(name="Calibri", color="FFFFFF", bold=True, size=20)
    cover["A1"].fill = cover_fill
    cover["A1"].alignment = Alignment(wrap_text=True, vertical="center", indent=2)
    cover.row_dimensions[1].height = 60
    cover["A2"] = f"Client: {client_name}" if client_name else "SAP S/4HANA Cloud — Public Edition"
    cover["A2"].font = Font(name="Calibri", color="DDEEFF", size=13)
    cover["A2"].fill = PatternFill("solid", fgColor="2563EB")
    cover["A2"].alignment = Alignment(vertical="center", indent=2)
    cover.row_dimensions[2].height = 32
    if overview_text:
        cover["A3"] = overview_text
        cover["A3"].font = Font(name="Calibri", color="1F2937", size=11)
        cover["A3"].alignment = Alignment(wrap_text=True, vertical="top", indent=2)
        cover.row_dimensions[3].height = max(80, len(overview_text) // 3)

    try:
        from openpyxl.drawing.image import Image as XLImage
        img_row = 4 if overview_text else 3
        for b64_data, col_anchor in [(client_logo_b64, "A"), (partner_logo_b64, "C")]:
            if not b64_data:
                continue
            raw = b64_data.split(",", 1)[-1]
            xl_img = XLImage(io.BytesIO(base64.b64decode(raw)))
            xl_img.width, xl_img.height = 160, 80
            cover.add_image(xl_img, f"{col_anchor}{img_row}")
    except Exception:
        pass

    # Data sheet
    ws = wb.create_sheet("BDCQ Questions")
    header_fill2 = PatternFill("solid", fgColor="1F4E79")
    ws.append(columns)
    for cell in ws[1]:
        cell.font  = Font(color="FFFFFF", bold=True, size=9)
        cell.fill  = header_fill2
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for row in all_rows:
        ws.append([str(row.get(col) or "") for col in columns])
        for cell in ws[ws.max_row]:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    for i in range(1, len(columns) + 1):
        ws.column_dimensions[get_column_letter(i)].width = 38
    ws.freeze_panes = "A2"

    # ── Save to output/ ───────────────────────────────────────────────────────
    today = date.today().strftime("%Y%m%d")
    if not filename_stem:
        safe = re.sub(r"[^\w]", "_", client_name or "Client")[:20]
        filename_stem = f"{safe}_BDCQ_{today}"
    filename = re.sub(r"[^\w.\-]", "_", filename_stem) + ".xlsx"
    out_path = OUTPUT_DIR / filename
    wb.save(str(out_path))

    return {"ok": True, "filename": filename, "rowCount": len(all_rows)}
